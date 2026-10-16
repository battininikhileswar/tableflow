"""Real engine and HTTP integration tests. Run: python -m unittest discover -s tests -v"""
import io,json,os,shutil,socket,subprocess,sys,tempfile,time,unittest,urllib.request,urllib.error
from pathlib import Path
import fitz
from PIL import Image
from openpyxl import load_workbook
from app.engine import extract,export_xlsx,validate
ROOT=Path(__file__).resolve().parents[1]

class EngineTests(unittest.TestCase):
    def test_native_exact_cells_and_provenance(self):
        result=extract(ROOT/'samples/sample-inventory.pdf','sample.pdf')
        truth=json.loads((ROOT/'samples/ground-truth.json').read_text())
        self.assertEqual(len(result['tables']),1)
        t=result['tables'][0]
        self.assertEqual(t['headers'],truth['headers']);self.assertEqual(t['rows'],truth['rows'])
        self.assertEqual(t['row_pages'],[1]*8+[2]*8)
    @unittest.skipUnless(shutil.which('tesseract') or os.getenv('TESSERACT_CMD'), "Tesseract OCR binary not installed")
    def test_scanned_wide_tables(self):
        result=extract(ROOT/'samples/sample-scanned.pdf','scan.pdf')
        self.assertEqual(result['page_count'],2);self.assertEqual(len(result['tables']),1)
        t=result['tables'][0];self.assertEqual(len(t['headers']),12);self.assertEqual(len(t['rows']),16)
        truth=json.loads((ROOT/'samples/ground-truth.json').read_text())
        matched=sum(a==b for r,e in zip(t['rows'],truth['rows']) for a,b in zip(r,e))
        self.assertGreaterEqual(matched/192,.90)
        self.assertTrue(all(len(c)==12 for c in t['confidence']))
    @unittest.skipUnless(shutil.which('tesseract') or os.getenv('TESSERACT_CMD'), "Tesseract OCR binary not installed")
    def test_image_ocr(self):
        result=extract(ROOT/'samples/sample-scan.png','scan.png')
        self.assertEqual(len(result['tables'][0]['headers']),12)
        self.assertEqual(len(result['tables'][0]['rows']),8)
    def test_fifty_pages(self):
        result=extract(ROOT/'samples/sample-50-pages.pdf','50.pdf')
        self.assertEqual(result['page_count'],50)
        self.assertEqual(sum(len(t['rows']) for t in result['tables']),50)
    def test_reject_fifty_one(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'51.pdf';doc=fitz.open()
            for i in range(51):doc.new_page()
            doc.save(path);doc.close()
            with self.assertRaisesRegex(ValueError,'50 pages'):extract(path,'51.pdf')
    def test_locked_pdf(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'locked.pdf';doc=fitz.open();doc.new_page();doc.save(path,encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='secret');doc.close()
            with self.assertRaisesRegex(ValueError,'Password'):extract(path,'locked.pdf')
    def test_blank_page_preserved(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'blank.png';Image.new('RGB',(500,500),'white').save(path)
            r=extract(path,'blank.png');self.assertEqual(r['tables'],[]);self.assertTrue(r['notes'])
    def test_wrong_totals_and_arithmetic(self):
        table={'id':'table-1','headers':['Item','Qty','Unit price','Amount'],'rows':[['A','2','50','100'],['B','3','40','130'],['Grand total','','','999']]}
        issues=validate([table]);self.assertTrue(any(i['kind']=='arithmetic' and i['row']==1 for i in issues));self.assertTrue(any(i['kind']=='total' and i['row']==2 for i in issues))
        table['rows'][1][3]='120';table['rows'][2][3]='220'
        self.assertFalse(any(i['kind'] in ('arithmetic','total') for i in validate([table])))
    def test_numeric_column_misalignment(self):
        table={'id':'table-1','headers':['Item','Amount'],'rows':[['A','1'],['B','2'],['C','3'],['D','wrong-text']]}
        self.assertTrue(any(i['kind']=='alignment' and i['row']==3 for i in validate([table])))
    def test_business_pdf_fixtures_and_totals(self):
        for name in ['sample-invoice','sample-bank-statement']:
            result=extract(ROOT/'samples'/(name+'.pdf'),name+'.pdf');truth=json.loads((ROOT/'samples'/(name+'-ground-truth.json')).read_text())
            self.assertEqual(result['tables'][0]['headers'],truth['headers']);self.assertEqual(result['tables'][0]['rows'],truth['rows'])
            if name=='sample-invoice':self.assertTrue({'total','arithmetic'}<={i['kind'] for i in result['issues']})
    def test_rule_flags(self):
        t={'id':'table-1','headers':['A','B'],'rows':[['','x'],['','x']],'confidence':[[99,50],[99,90]],'method':'ocr-alignment'}
        kinds={i['kind'] for i in validate([t])}
        self.assertTrue({'empty','duplicate','ocr','alignment'}<=kinds)
    def test_export_preserves_ids_and_blocks_formulas(self):
        r=extract(ROOT/'samples/sample-inventory.pdf','=HYPERLINK("bad")')
        r['tables'][0]['rows'][0][0]='00123';r['tables'][0]['rows'][0][1]='=1+1'
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'out.xlsx';export_xlsx(r,p);wb=load_workbook(p)
            self.assertEqual(wb['Table 1']['A2'].value,'00123')
            self.assertEqual(wb['Table 1']['B2'].data_type,'s')
            self.assertEqual(wb['Overview']['B2'].data_type,'s')
            self.assertEqual(wb['Table 1'].freeze_panes,'A2');self.assertEqual(wb['Table 1'].max_column,13)
            self.assertIn('Validation',wb.sheetnames);self.assertIn('Page text',wb.sheetnames)

class APITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();sock=socket.socket();sock.bind(('127.0.0.1',0));cls.port=sock.getsockname()[1];sock.close()
        cls.base=f'http://127.0.0.1:{cls.port}'
        cls.env={**os.environ,'PORT':str(cls.port),'TABLEFLOW_DATA':cls.temp.name,'GEMINI_API_KEY':''}
        cls.start_server()
        req=urllib.request.Request(cls.base+'/api/auth/register',data=json.dumps({'name':'Test User','email':'judge@example.test','password':'Test-pass-2026'}).encode(),method='POST')
        with urllib.request.urlopen(req) as response:cls.cookie=response.headers['Set-Cookie'].split(';')[0]
    @classmethod
    def start_server(cls):
        cls.proc=subprocess.Popen([sys.executable,'run.py'],cwd=ROOT,env=cls.env,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        for i in range(100):
            try:urllib.request.urlopen(cls.base+'/api/health',timeout=1).close();return
            except Exception:time.sleep(.05)
        raise RuntimeError('Server did not start')
    @classmethod
    def tearDownClass(cls):
        cls.proc.terminate();cls.proc.wait(timeout=10);cls.temp.cleanup()
    def request(self,path,method='GET',body=None,headers=None):
        req=urllib.request.Request(self.base+path,data=body,method=method,headers={'Cookie':self.cookie,**(headers or {})})
        try:
            with urllib.request.urlopen(req,timeout=30) as r:return r.status,r.read(),dict(r.headers)
        except urllib.error.HTTPError as e:return e.code,e.read(),dict(e.headers)
    def upload(self,content=None,name='sample.pdf'):
        code,body,_=self.request('/api/jobs?mode=local&name='+name,'POST',content or (ROOT/'samples/sample-inventory.pdf').read_bytes())
        self.assertEqual(code,202,body);return json.loads(body)['id']
    def wait(self,jid):
        for _ in range(200):
            job=json.loads(self.request('/api/jobs/'+jid)[1])
            if job['status'] in ('complete','failed'):return job
            time.sleep(.05)
        self.fail('Job timed out')
    def test_full_upload_edit_export_retry_delete(self):
        jid=self.upload();job=self.wait(jid);self.assertEqual(job['status'],'complete')
        tables=job['result']['tables'];tables[0]['rows'][0][0]='EDITED'
        code,body,_=self.request('/api/jobs/'+jid+'/tables','PUT',json.dumps({'tables':tables}).encode())
        self.assertEqual(code,200,body)
        code,body,headers=self.request('/api/jobs/'+jid+'/export');self.assertEqual(code,200)
        wb=load_workbook(io.BytesIO(body));self.assertEqual(wb['Table 1']['A2'].value,'EDITED')
        self.assertTrue(headers['Content-Type'].endswith('spreadsheetml.sheet'))
        code,body,_=self.request('/api/jobs/'+jid+'/page/2.png');self.assertEqual(code,200);self.assertTrue(body.startswith(b'\x89PNG'))
        code,body,_=self.request('/api/jobs/'+jid+'/report');self.assertEqual(code,200);self.assertEqual(json.loads(body)['filename'],'sample.pdf')
        code,body,_=self.request('/api/jobs/'+jid+'/retry','POST');self.assertEqual(code,202)
        self.assertEqual(self.wait(json.loads(body)['id'])['status'],'complete')
        self.assertEqual(self.request('/api/jobs/'+jid,'DELETE')[0],200)
        self.assertEqual(self.request('/api/jobs/'+jid)[0],404)
    def test_preview_while_another_job_extracts(self):
        first=self.upload();self.wait(first)
        scan=self.upload((ROOT/'samples/sample-scanned.pdf').read_bytes(),'scan.pdf')
        code,body,_=self.request('/api/jobs/'+first+'/page/1.png')
        self.assertEqual(code,200);self.assertTrue(body.startswith(b'\x89PNG'))
        self.assertEqual(self.wait(scan)['status'],'complete')
    def test_accounts_require_session_and_isolate_documents(self):
        self.assertEqual(self.request('/api/jobs',headers={'Cookie':''})[0],401)
        jid=self.upload();self.wait(jid)
        code,body,headers=self.request('/api/auth/register','POST',json.dumps({'name':'Other User','email':'other@example.test','password':'Another-pass-2026'}).encode(),{'Cookie':''})
        self.assertEqual(code,200,body);other=headers['Set-Cookie'].split(';')[0]
        self.assertEqual(self.request('/api/jobs/'+jid,headers={'Cookie':other})[0],404)
        self.assertEqual(json.loads(self.request('/api/jobs',headers={'Cookie':other})[1]),[])
        self.assertEqual(self.request('/api/jobs/'+jid,'DELETE',headers={'Cookie':other})[0],404)
        self.assertEqual(self.request('/api/auth/login','POST',b'{"email":"other@example.test","password":"wrongpass"}')[0],400)
        code,_,headers=self.request('/api/auth/login','POST',b'{"email":"other@example.test","password":"Another-pass-2026"}')
        self.assertEqual(code,200)
        session=headers['Set-Cookie'].split(';')[0]
        self.assertIn('HttpOnly',headers['Set-Cookie']);self.assertIn('SameSite=Strict',headers['Set-Cookie'])
        self.assertEqual(self.request('/api/auth/logout','POST',headers={'Cookie':session})[0],200)
        self.assertIsNone(json.loads(self.request('/api/auth/me',headers={'Cookie':session})[1])['user'])
    def test_plan_preference_and_invalid_plan(self):
        code,body,_=self.request('/api/plan','PUT',b'{"plan":"pro"}')
        self.assertEqual(code,200);self.assertEqual(json.loads(body)['user']['plan'],'pro')
        self.assertEqual(self.request('/api/plan','PUT',b'{"plan":"unknown"}')[0],400)
    def test_share_snapshot_survives_edits_and_can_be_revoked(self):
        jid=self.upload();job=self.wait(jid)
        code,body,_=self.request('/api/jobs/'+jid+'/share','POST',b'{"hours":24}')
        self.assertEqual(code,201,body);share=json.loads(body)
        code,content,_=self.request(share['path'],headers={'Cookie':''});self.assertEqual(code,200)
        original=load_workbook(io.BytesIO(content))['Table 1']['A2'].value
        job['result']['tables'][0]['rows'][0][0]='CHANGED-AFTER-SHARE'
        self.assertEqual(self.request('/api/jobs/'+jid+'/tables','PUT',json.dumps({'tables':job['result']['tables']}).encode())[0],200)
        self.assertEqual(load_workbook(io.BytesIO(self.request(share['path'])[1]))['Table 1']['A2'].value,original)
        self.assertEqual(self.request('/api/jobs/'+jid+'/share','POST',b'{"hours":999}')[0],400)
        self.assertEqual(self.request('/api/shares/'+share['id'],'DELETE')[0],200)
        self.assertEqual(self.request(share['path'],headers={'Cookie':''})[0],410)
    def test_deleting_document_revokes_its_shares(self):
        jid=self.upload();self.wait(jid)
        share=json.loads(self.request('/api/jobs/'+jid+'/share','POST',b'{"hours":1}')[1])
        self.assertEqual(self.request('/api/jobs/'+jid,'DELETE')[0],200)
        self.assertEqual(self.request(share['path'])[0],410)
    def test_no_visual_image_elements_in_ui(self):
        html=self.request('/')[1].decode()
        self.assertNotIn('<img',html.lower());self.assertNotIn('<svg',html.lower())
    def test_expiring_share_and_original_download(self):
        import sqlite3
        jid=self.upload();self.wait(jid)
        self.assertTrue(self.request('/api/jobs/'+jid+'/source')[1].startswith(b'%PDF-'))
        share=json.loads(self.request('/api/jobs/'+jid+'/share','POST',b'{"hours":1}')[1])
        with sqlite3.connect(Path(self.temp.name)/'accounts.sqlite3') as db:db.execute('UPDATE shares SET expires=0 WHERE id=?',(share['id'],))
        self.assertEqual(self.request(share['path'],headers={'Cookie':''})[0],410)
    def test_json_object_validation(self):
        self.assertEqual(self.request('/api/plan','PUT',b'[]')[0],400)
        self.assertEqual(self.request('/api/auth/login','POST',b'[]')[0],400)
    def test_ai_settings_are_private_and_missing_key_is_actionable(self):
        code,body,_=self.request('/api/ai/settings');self.assertEqual(code,200);self.assertFalse(json.loads(body)['configured'])
        code,body,_=self.request('/api/jobs?mode=ai&name=test.pdf','POST',(ROOT/'samples/sample-inventory.pdf').read_bytes())
        self.assertEqual(code,400);self.assertIn('Gemini API key',json.loads(body)['error'])
        secret='fixture-key-do-not-use-live'
        code,body,_=self.request('/api/ai/settings','PUT',json.dumps({'model':'gemini-3.8-flash','api_key':secret}).encode())
        self.assertEqual(code,200);self.assertTrue(json.loads(body)['configured']);self.assertNotIn(secret,body.decode())
        self.assertNotIn(secret,(Path(self.temp.name)/'accounts.sqlite3').read_bytes().decode(errors='ignore'))
        self.assertEqual(self.request('/api/ai/settings','PUT',b'{"clear":true}')[0],200)
    def test_invalid_upload(self):
        code,body,_=self.request('/api/jobs?name=bad.txt','POST',b'bad data');self.assertEqual(code,400)
    def test_corrupt_pdf_becomes_failed_job(self):
        jid=self.upload(b'%PDF-corrupt');self.assertEqual(self.wait(jid)['status'],'failed')
    def test_origin_and_host_guards(self):
        self.assertEqual(self.request('/api/jobs','POST',b'x',{'Origin':'https://example.org'})[0],403)
        self.assertEqual(self.request('/api/health',headers={'Host':'evil.example'})[0],403)
    def test_unknown_routes(self):
        self.assertEqual(self.request('/api/not-real')[0],404)
        self.assertEqual(self.request('/../requirements.txt')[0],404)
        self.assertEqual(self.request('/api/jobs/'+'a'*32)[0],404)
    def test_invalid_edit_is_rejected(self):
        jid=self.upload();self.wait(jid)
        self.assertEqual(self.request('/api/jobs/'+jid+'/tables','PUT',b'{"tables":[]}')[0],400)
        self.assertEqual(self.request('/api/jobs/'+jid+'/page/99.png')[0],400)
    def test_static_assets(self):
        for path in ['/','/style.css','/app.js','/product.css','/product.js','/samples/sample-inventory.pdf']:
            code,body,headers=self.request(path);self.assertEqual(code,200,path);self.assertTrue(body)
    def test_persistence_across_restart(self):
        jid=self.upload();self.wait(jid)
        self.proc.terminate();self.proc.wait(timeout=10);self.__class__.start_server()
        job=json.loads(self.request('/api/jobs/'+jid)[1]);self.assertEqual(job['status'],'complete')

if __name__=='__main__':unittest.main()
