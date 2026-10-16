"""AI contracts and orchestration use deterministic fakes, NOT live AI accuracy."""
import copy,io,json,tempfile,unittest,urllib.error
from pathlib import Path
from app.ai_agent import GeminiClient,AIError,AIProviderError,check_output,extract_ai,revalidate
from app.engine import export_xlsx
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]

def output(rows=None,continuation='',uncertain=None):
    return {'tables':[{'title':'Items','headers':['Item','Qty','Price','Amount'],'rows':rows or [['A','2','10','20']],'continues_table':continuation,'uncertain_cells':uncertain or []}],'warnings':[]}
class Reply:
    def __init__(self,data):self.data=json.dumps(data).encode()
    def __enter__(self):return self
    def __exit__(self,*args):pass
    def read(self,*args):return self.data

def wire(out,finish='STOP'):
    return {'candidates':[{'finishReason':finish,'content':{'parts':[{'text':json.dumps(out)}]}}]}
class FakeClient:
    def __init__(self,replies):self.replies=iter(replies);self.model='fixture-model';self.calls=0
    def generate(self,*args):
        self.calls+=1;item=next(self.replies)
        if isinstance(item,Exception):raise item
        return copy.deepcopy(item)
class AIContractTests(unittest.TestCase):
    def test_real_request_contract_and_no_key_in_url(self):
        captured=[]
        def transport(req,timeout):captured.append(req);return Reply(wire(output()))
        client=GeminiClient('secret-fixture',transport=transport)
        self.assertEqual(client.generate(b'png','extract')['tables'][0]['rows'][0][0],'A')
        body=json.loads(captured[0].data);self.assertIn('responseFormat',body['generationConfig'])
        self.assertEqual(body['contents'][0]['parts'][1]['inlineData']['mimeType'],'image/png')
        self.assertNotIn('secret-fixture',captured[0].full_url)
    def test_misalignment_and_numeric_json_rejected(self):
        obj=output();obj['tables'][0]['rows']=[['too','short']]
        with self.assertRaises(AIError):check_output(obj)
        obj=output();obj['tables'][0]['rows'][0][1]=2
        with self.assertRaises(AIError):check_output(obj)
    def test_truncated_result_retried_then_rejected(self):
        client=GeminiClient('fixture',transport=lambda *a,**k:Reply(wire(output(),'MAX_TOKENS')),sleeper=lambda n:None)
        with self.assertRaises(AIError):client.generate(b'x','x')
        self.assertEqual(client.calls,3)
    def test_rate_limit_retries(self):
        count=[0]
        def transport(req,timeout):
            count[0]+=1
            if count[0]<3:raise urllib.error.HTTPError(req.full_url,429,'rate',{'Retry-After':'1'},None)
            return Reply(wire(output()))
        client=GeminiClient('fixture',transport=transport,sleeper=lambda n:None)
        self.assertTrue(client.generate(b'x','x')['tables']);self.assertEqual(client.calls,3)
    def test_credentials_fail_fast(self):
        def transport(req,timeout):raise urllib.error.HTTPError(req.full_url,403,'denied',{},None)
        client=GeminiClient('fixture',transport=transport,sleeper=lambda n:None)
        with self.assertRaises(AIProviderError):client.generate(b'x','x')
        self.assertEqual(client.calls,1)
    def test_two_pass_merge_audit_uncertainty_and_excel(self):
        first=output([['A','2','10','2O']]);verified=output([['A','2','10','20']])
        second=output([['B','3','10','30']],continuation='table-1',uncertain=[{'row':0,'col':0,'reason':'Faint label'}])
        client=FakeClient([first,verified,second,second])
        r=extract_ai(ROOT/'samples/sample-inventory.pdf','sample.pdf',client=client)
        self.assertEqual(client.calls,4);self.assertEqual(r['ai']['verified_pages'],2)
        self.assertEqual(len(r['tables']),1);self.assertEqual(r['tables'][0]['row_pages'],[1,2])
        self.assertEqual(r['tables'][0]['rows'][0][3],'20');self.assertTrue(r['ai_audit'])
        self.assertTrue(any(i['kind']=='ai_uncertain' and i['row']==1 for i in r['issues']))
        r['tables'][0]['rows'][1][0]='edited';self.assertTrue(any(i['kind']=='ai_uncertain' for i in revalidate(r)))
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'ai.xlsx';export_xlsx(r,path);wb=load_workbook(path)
            self.assertIn('AI audit',wb.sheetnames);self.assertEqual(wb['Table 1']['D2'].value,'20')
    def test_verification_failure_does_not_accept_draft(self):
        client=FakeClient([output(),AIError('verification failed'),output(),output()])
        r=extract_ai(ROOT/'samples/sample-inventory.pdf','sample.pdf',client=client)
        self.assertEqual(r['ai']['failed_pages'],1);self.assertTrue(r['ai']['partial'])
        self.assertEqual(r['tables'][0]['row_pages'],[2])
    def test_no_successful_page_is_failure(self):
        with self.assertRaisesRegex(AIError,'any pages'):
            extract_ai(ROOT/'samples/sample-inventory.pdf','sample.pdf',client=FakeClient([AIError('offline'),AIError('offline')]))
    def test_incompatible_continuation_kept_separate(self):
        first=output();second=output(continuation='table-1');second['tables'][0]['headers']=['Other','Qty','Price','Amount']
        r=extract_ai(ROOT/'samples/sample-inventory.pdf','sample.pdf',client=FakeClient([first,first,second,second]))
        self.assertEqual(len(r['tables']),2);self.assertTrue(any('kept separate' in n for n in r['notes']))
    def test_printed_arithmetic_error_preserved_and_flagged(self):
        wrong=output([['A','2','10','99']])
        r=extract_ai(ROOT/'samples/sample-scan.png','image.png',client=FakeClient([wrong,wrong]))
        self.assertEqual(r['tables'][0]['rows'][0][3],'99');self.assertTrue(any(i['kind']=='arithmetic' for i in r['issues']))

class AIServerOrchestrationTests(unittest.TestCase):
    def test_ai_job_automatically_writes_workbook_and_keeps_key_out_of_state(self):
        import os,time
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as tmp,patch.dict(os.environ,{'TABLEFLOW_DATA':tmp,'GEMINI_API_KEY':'private-fixture-key'}):
            from app import server
            page=output()
            result=extract_ai(ROOT/'samples/sample-scan.png','scan.png',client=FakeClient([page,page]))
            with patch.object(server,'extract_ai',return_value=result) as agent:
                jid=server.create_job((ROOT/'samples/sample-scan.png').read_bytes(),'scan.png',owner='fixture-owner',mode='ai')
                for _ in range(200):
                    if server.JOBS[jid]['status'] in ('complete','failed'):break
                    time.sleep(.02)
                job=server.JOBS[jid]
                self.assertEqual(job['status'],'complete',job.get('error'));self.assertTrue(job['workbook_ready'])
                self.assertTrue((server.DATA/jid/'extracted.xlsx').is_file())
                self.assertNotIn('private-fixture-key',(server.DATA/jid/'job.json').read_text())
                self.assertEqual(agent.call_args.kwargs['key'],'private-fixture-key')
